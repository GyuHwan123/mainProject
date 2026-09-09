"""Render the same document sheets used by the finance preview as printable PDFs."""
from html import escape
from io import BytesIO

from openpyxl import load_workbook
from reportlab.lib import colors
from reportlab.lib.enums import TA_CENTER, TA_RIGHT
from reportlab.lib.pagesizes import A4, landscape
from reportlab.lib.styles import ParagraphStyle
from reportlab.lib.units import mm
from reportlab.platypus import PageBreak, Paragraph, SimpleDocTemplate, Spacer, Table, TableStyle

from app.services.finance_workbook_service import build_finance_workbook, SHEET_NAMES
from app.services.knowledge_report_service import _font


def build_finance_pdf(records: list[dict], *, author: dict | None = None) -> bytes:
    workbook = load_workbook(BytesIO(build_finance_workbook(records, author=author)))
    output = BytesIO()
    font = _font()
    if font == "Helvetica":
        raise RuntimeError("PDF 생성에 필요한 한글 글꼴이 없습니다.")
    ink = colors.HexColor("#173765")
    border = colors.HexColor("#cbd5e1")
    shade = colors.HexColor("#f1f5f9")
    width = landscape(A4)[0] - 32 * mm
    style = ParagraphStyle("FinanceCell", fontName=font, fontSize=8, leading=12, textColor=ink, wordWrap="CJK")

    def cell(value, align=0, size=8):
        text = f"{value:,.3f}".rstrip("0").rstrip(".") if isinstance(value, (int, float)) else str(value or "")
        return Paragraph(escape(text).replace("\n", "<br/>"), ParagraphStyle("Cell", parent=style, alignment=align, fontSize=size, leading=size * 1.5))

    def table(rows, widths, repeat=0):
        result = Table(rows, colWidths=widths, repeatRows=repeat, hAlign="LEFT")
        result.setStyle(TableStyle([
            ("GRID", (0, 0), (-1, -1), .5, border),
            ("VALIGN", (0, 0), (-1, -1), "MIDDLE"),
            ("TOPPADDING", (0, 0), (-1, -1), 9),
            ("BOTTOMPADDING", (0, 0), (-1, -1), 9),
            ("LEFTPADDING", (0, 0), (-1, -1), 7),
            ("RIGHTPADDING", (0, 0), (-1, -1), 7),
        ]))
        return result

    story = []
    try:
        for name in SHEET_NAMES.values():
            if name not in workbook.sheetnames:
                continue
            if story:
                story.append(PageBreak())
            rows = list(workbook[name].values)
            approval = table([[cell(label, TA_CENTER) for label in ["기안자", "검토자", "승인자"]],
                              [cell(rows[2][column], TA_CENTER) for column in [5, 6, 7]]], [25 * mm] * 3)
            approval.setStyle(TableStyle([("BACKGROUND", (0, 0), (-1, 0), shade)]))
            heading = Table([[cell(rows[1][0] or name, size=20), approval]], colWidths=[width - 75 * mm, 75 * mm])
            heading.setStyle(TableStyle([("VALIGN", (0, 0), (-1, -1), "MIDDLE"), ("LEFTPADDING", (0, 0), (-1, -1), 0), ("RIGHTPADDING", (0, 0), (-1, -1), 0)]))
            story.extend([heading, Spacer(1, 8 * mm)])
            metadata = table([[cell(row[0]), cell(row[1], TA_RIGHT), cell(row[4]), cell(row[5], TA_RIGHT)] for row in rows[4:8]], [width * fraction for fraction in [.11, .39, .11, .39]])
            metadata.setStyle(TableStyle([("BACKGROUND", (0, 0), (0, -1), shade), ("BACKGROUND", (2, 0), (2, -1), shade)]))
            story.extend([metadata, Spacer(1, 7 * mm), cell("지출 내역", size=11), Spacer(1, 3 * mm)])
            headers = list(rows[10])
            while headers and headers[-1] is None:
                headers.pop()
            details = [row[:len(headers)] for row in rows[11:] if any(value is not None and value != "" for value in row) and not str(row[0] or "").startswith("품목금액 합계")]
            weights = {"영수증 ID": .9, "품목 순번": .65, "지출 카테고리": 1.2, "거래일": 1.05, "상호명": 1.9, "품목명": 1.7, "수량": .6, "단가": .85, "품목금액": 1}
            column_weights = [weights.get(label, 1) for label in headers]
            rendered = [[cell(label, TA_CENTER) for label in headers]]
            for row in details:
                rendered.append([cell((str(value)[:8] + "…") if index == 0 and len(str(value or "")) > 12 else value,
                                      TA_RIGHT if isinstance(value, (int, float)) else TA_CENTER)
                                 for index, value in enumerate(row)])
            amounts = [row[-1] for row in details if isinstance(row[-1], (int, float))]
            rendered.append([cell("품목금액 합계 (미확인 금액 제외)")] + [cell("")] * (len(headers) - 2) + [cell(f"{sum(amounts):,} 원" if amounts else "—", TA_RIGHT)])
            detail_table = table(rendered, [width * weight / sum(column_weights) for weight in column_weights], repeat=1)
            detail_table.setStyle(TableStyle([("BACKGROUND", (0, 0), (-1, 0), shade), ("ROWBACKGROUNDS", (0, 1), (-1, -2), [colors.white, colors.HexColor("#f8fafc")]), ("SPAN", (0, -1), (-2, -1)), ("BACKGROUND", (0, -1), (-1, -1), shade)]))
            story.append(detail_table)
        document = SimpleDocTemplate(output, pagesize=landscape(A4), leftMargin=16 * mm, rightMargin=16 * mm, topMargin=14 * mm, bottomMargin=14 * mm, title="재무 문서", author=(author or {}).get("name", ""))
        def footer(canvas, doc):
            canvas.setFont(font, 8)
            canvas.setFillColor(ink)
            canvas.drawRightString(landscape(A4)[0] - 16 * mm, 8 * mm, str(doc.page))
        document.build(story, onFirstPage=footer, onLaterPages=footer)
        return output.getvalue()
    finally:
        workbook.close()
